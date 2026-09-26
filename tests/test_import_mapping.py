import base64
import csv
import importlib.util
import io
import json
import os
import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from api import import_mapping


class MappingConversionTest(unittest.TestCase):
    def test_agent_returns_a_structured_mapping_proposal(self):
        content = b"Date,Dish,Sold,Covers\n2025-09-25,Soup,4,30\n"
        details, fallback = import_mapping.describe("sales.csv", content)
        response = {"output": [{"content": [{"type": "output_text", "text": json.dumps(fallback)}]}]}
        with patch.dict(os.environ, {"YLD_IMPORT_AGENT_API_KEY": "test-key"}), patch.object(import_mapping, "urlopen") as opened:
            opened.return_value.__enter__.return_value = io.BytesIO(json.dumps(response).encode())
            self.assertEqual(import_mapping.agent_mapping(details, fallback), fallback)
            request = json.loads(opened.call_args.args[0].data)
            self.assertEqual(request["text"]["format"]["type"], "json_schema")
            self.assertFalse(request["store"])

    def test_transaction_csv_maps_dates_and_sums_repeated_dishes(self):
        content = ("Trading Day;Menu Item;Qty;Pax;Made\n"
                   "25/09/2025;Soup;2;30;3\n"
                   "25/09/2025;Soup;4;30;5\n").encode()
        details, mapping = import_mapping.describe("sales.csv", content)
        self.assertEqual(details["sheets"][0]["name"], "CSV")
        self.assertEqual(mapping["columns"]["sold"], "Qty")
        self.assertEqual(mapping["columns"]["covers"], "Pax")
        mapping["date_format"] = "dmy"
        mapping["aggregation"] = "sum"
        rows = list(csv.DictReader(io.StringIO(import_mapping.normalize("sales.csv", content, mapping))))
        self.assertEqual(len(rows), 1)
        self.assertEqual((rows[0]["date"], rows[0]["dish"], rows[0]["sold"], rows[0]["prepared"]),
                         ("2025-09-25", "Soup", "6", "8"))

    def test_wide_xlsx_with_title_row_and_excel_dates(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Service sales"
        sheet.append(["Service report"])
        sheet.append(["Date", "Guests", "Soup", "Steak"])
        sheet.append([date(2025, 9, 25), 30, 4, 8])
        sheet.append([date(2025, 9, 26), 31, 5, None])
        output = io.BytesIO()
        workbook.save(output)
        content = output.getvalue()
        details, mapping = import_mapping.describe("report.xlsx", content)
        self.assertEqual(details["sheets"][0]["header_row"], 2)
        self.assertEqual(mapping["layout"], "wide")
        self.assertEqual(set(mapping["dish_columns"]), {"Soup", "Steak"})
        rows = list(csv.DictReader(io.StringIO(import_mapping.normalize("report.xlsx", content, mapping))))
        self.assertEqual([(row["date"], row["dish"], row["sold"]) for row in rows],
                         [("2025-09-25", "Soup", "4"), ("2025-09-25", "Steak", "8"), ("2025-09-26", "Soup", "5")])

    def test_menu_sheet_maps_item_cost_and_price(self):
        content = b"Menu item,Food cost,Selling price,Course\nSoup,3.50,14.00,Starters\n"
        details, mapping = import_mapping.describe_menu("menu.csv", content)
        self.assertEqual(details["sheets"][0]["name"], "CSV")
        self.assertEqual(mapping["columns"]["dish"], "Menu item")
        menu = import_mapping.normalize_menu("menu.csv", content, mapping)
        self.assertEqual(menu, [{"name": "Soup", "ingredient_cost": "3.50", "price": "14.00", "category": "Starters"}])

    def test_menu_header_detection_ignores_sales_report_columns(self):
        workbook = Workbook()
        report = workbook.active
        report.title = "Sales"
        report.append(["Date", "Dish", "Sold", "Covers"])
        report.append(["2025-09-25", "Soup", 4, 30])
        menu = workbook.create_sheet("Menu")
        menu.append(["Menu export"])
        menu.append(["Menu item", "Food cost", "Selling price"])
        menu.append(["Soup", 3.5, 14])
        output = io.BytesIO()
        workbook.save(output)
        _, mapping = import_mapping.describe_menu("kitchen.xlsx", output.getvalue())
        self.assertEqual((mapping["sheet"], mapping["header_row"]), ("Menu", 2))

    def test_missing_covers_and_unknown_headers_are_rejected(self):
        content = b"Date,Item,Qty\n2025-09-25,Soup,4\n"
        _, mapping = import_mapping.describe("sales.csv", content)
        with self.assertRaisesRegex(ValueError, "covers"):
            import_mapping.normalize("sales.csv", content, mapping)
        mapping["columns"]["covers"] = "invented"
        with self.assertRaisesRegex(ValueError, "Mapped columns"):
            import_mapping.normalize("sales.csv", content, mapping)

    def test_owner_can_select_a_header_row_when_detection_picks_a_report_title(self):
        content = ("Date,Dish,Sold,Covers,Report notes\n"
                   "When,What,HowMany,Guests\n"
                   "2025-09-25,Soup,4,30\n").encode()
        details, mapping = import_mapping.describe("sales.csv", content)
        self.assertEqual(mapping["header_row"], 1)
        options = details["sheets"][0]["header_options"]
        self.assertEqual(options[1]["headers"], ["When", "What", "HowMany", "Guests"])
        mapping["header_row"] = 2
        mapping["columns"].update(date="When", dish="What", sold="HowMany", covers="Guests")
        rows = list(csv.DictReader(io.StringIO(import_mapping.normalize("sales.csv", content, mapping))))
        self.assertEqual((rows[0]["date"], rows[0]["dish"], rows[0]["sold"], rows[0]["covers"]),
                         ("2025-09-25", "Soup", "4", "30"))


class MappingPersistenceTest(unittest.TestCase):
    def test_confirmed_mapping_is_reused_for_same_workspace(self):
        with tempfile.TemporaryDirectory() as directory, patch.dict(os.environ, {"YLD_DB_PATH": str(Path(directory) / "test.db"), "YLD_IMPORT_AGENT_API_KEY": ""}):
            path = Path(__file__).resolve().parents[1] / "api" / "main.py"
            spec = importlib.util.spec_from_file_location("yld_import_mapping_test", path)
            module = importlib.util.module_from_spec(spec)
            sys.modules[spec.name] = module
            try:
                spec.loader.exec_module(module)
                with module.db() as conn:
                    conn.execute("INSERT INTO workspaces(id,name,created_at,data_mode) VALUES (?,?,?,?)", ("workspace-one", "Kitchen", 1, "empty"))
                user = {"workspace_id": "workspace-one", "data_mode": "empty", "role": "owner"}
                content = b"Trading Day,Menu Item,Qty,Pax\n2025-09-25,Soup,4,30\n"
                payload = module.SpreadsheetInput(filename="sales.csv", content_base64=base64.b64encode(content).decode())
                with patch.object(import_mapping, "agent_mapping") as agent:
                    inspected = module.inspect_spreadsheet(payload, user)
                    agent.assert_not_called()
                self.assertEqual(inspected["mapping_source"], "rules")
                with patch.object(import_mapping, "agent_mapping", return_value=inspected["mapping"]) as agent:
                    opted_in = module.inspect_spreadsheet(module.SpreadsheetInput(**{**payload.model_dump(), "use_agent": True}), user)
                    agent.assert_called_once()
                    self.assertEqual(opted_in["mapping_source"], "agent")
                mapped = module.MappedSpreadsheetInput(**payload.model_dump(), mapping=inspected["mapping"])
                normalized = module.normalize_spreadsheet(mapped, user)
                self.assertEqual(normalized["preview"]["rows"], 1)
                dish_id = normalized["preview"]["dish_costs"][0]["id"]
                module.commit_import(module.SalesCsvInput(csv_text=normalized["csv_text"],
                    menu_costs={dish_id: module.DishCostsInput(ingredient_cost=3, price=12)},
                    source_mapping=inspected["mapping"], source_signature=inspected["signature"]), user)
                again = module.inspect_spreadsheet(payload, user)
                self.assertEqual(again["mapping_source"], "saved")
                with module.db() as conn:
                    conn.execute("INSERT INTO workspaces(id,name,created_at,data_mode) VALUES (?,?,?,?)", ("workspace-two", "Other Kitchen", 1, "empty"))
                other = module.inspect_spreadsheet(payload, {"workspace_id": "workspace-two", "data_mode": "empty", "role": "owner"})
                self.assertEqual(other["mapping_source"], "rules")
            finally:
                sys.modules.pop(spec.name, None)


if __name__ == "__main__":
    unittest.main()
